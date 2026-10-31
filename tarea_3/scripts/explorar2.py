from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

libro = load_workbook("datos/originales/Cap_12_ENDES_2025.xlsx", data_only=True)
hoja = libro["12.1"]

for col in range(2, 44):
    anio = hoja.cell(7, col).value
    detalle = hoja.cell(8, col).value
    if anio or detalle:
        print(get_column_letter(col), col, anio, "|", detalle)

print("Apurímac 2025:", hoja["AH63"].value, "IC:", hoja["AJ63"].value, "a", hoja["AK63"].value)