from openpyxl import load_workbook

libro = load_workbook("datos/originales/Cap_12_ENDES_2025.xlsx", data_only=True)
print(libro.sheetnames[:6])

hoja = libro["12.1"]
print("Título:", hoja["A2"].value)
print("Fila 11:", hoja["A11"].value, hoja["B11"].value)

for fila in range(58, 66):
    print(fila, hoja.cell(fila, 1).value)