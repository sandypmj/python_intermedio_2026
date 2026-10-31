import re
import pandas as pd
from openpyxl import load_workbook

def a_numero(valor):
    """54.86 -> 54.86 ;  "(7,7)" -> 7.7 ;  "¬" o vacío -> None"""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().strip("()").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None

libro = load_workbook("datos/originales/Cap_12_ENDES_2025.xlsx", data_only=True)
hoja = libro["12.4"]

# 1. ¿En qué columna está el valor de cada año?
columna_anio = {}
for col in range(2, 44):
    texto = str(hoja.cell(7, col).value).strip()
    if texto.isdigit():
        columna_anio[int(texto)] = col
print("Columnas:", columna_anio)

# 2. ¿En qué fila empiezan los departamentos?
for fila in range(1, hoja.max_row + 1):
    if hoja.cell(fila, 1).value == "Departamento":
        primera = fila + 1
        break

# 3. Leer fila por fila hasta llegar a una celda vacía
registros = []
fila = primera
while hoja.cell(fila, 1).value is not None:
    nombre = hoja.cell(fila, 1).value
    nombre = re.sub(r"\s+\d+/$", "", nombre)
    registro = {"departamento": nombre}
    for anio, col in columna_anio.items():
        registro[anio] = round(a_numero(hoja.cell(fila, col).value), 1)
    registros.append(registro)
    fila += 1

tabla = pd.DataFrame(registros)
tabla = tabla.sort_values(2025, ascending=False)
print(tabla.to_string(index=False))
print("Territorios:", len(tabla))
tabla.to_csv("datos/procesados/departamentos_fis_sex_12m.csv", index=False)